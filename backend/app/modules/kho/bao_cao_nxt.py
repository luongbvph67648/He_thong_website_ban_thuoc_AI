"""Báo cáo nhập - xuất - tồn (KHO13, spec 06 mục 3.19), tính hoàn toàn từ thẻ kho `bien_dong_kho`.

Tồn cuối = tồn đầu + tổng biến động trong kỳ (số có dấu: nhập dương, xuất âm). Ngày tính theo giờ Việt Nam.
"""

from datetime import date, timedelta
from enum import StrEnum
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy import Select, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.loi import LoiNghiepVu
from app.core.phan_trang import ThamSoPhanTrang, Trang, lay_trang, tao_trang
from app.core.thoi_gian import dau_ngay_vn
from app.core.tim_kiem import mau_tim_like
from app.models.enums import KhuTon, LoaiBienDong
from app.models.kho import BienDongKho, LoThuoc
from app.models.san_pham import SanPham

TOI_DA_DONG_XUAT_EXCEL = 20_000
TOI_DA_SO_NGAY = 3660
LOAI_THEO_COT = [(loai.value, loai) for loai in LoaiBienDong]  # thứ tự cột = thứ tự enum


class Nhom(StrEnum):
    san_pham = "san_pham"
    lo = "lo"


class DongNxt(BaseModel):
    san_pham_id: int
    sku: str
    ten: str
    lo_id: int | None = None
    so_lo: str | None = None
    han_su_dung: date | None = None
    ton_dau: int
    nhap_kho: int
    hoan_nhap: int
    nhan_tra_ve: int
    dieu_chinh: int
    xuat_ban: int
    huy_hang: int
    xuat_tra_ncc: int
    ton_cuoi: int


def _kiem_tra_ky(tu_ngay: date, den_ngay: date) -> None:
    loi = None
    if den_ngay < tu_ngay:
        loi = ("den_ngay", "Phải từ ngày bắt đầu trở đi")
    elif (den_ngay - tu_ngay).days > TOI_DA_SO_NGAY:
        loi = ("den_ngay", f"Kỳ báo cáo tối đa {TOI_DA_SO_NGAY} ngày")
    if loi:
        raise LoiNghiepVu("DU_LIEU_KHONG_HOP_LE", "Dữ liệu không hợp lệ", chi_tiet=[{"truong": loi[0], "loi": loi[1]}])


def _truy_van(
    *, tu_ngay: date, den_ngay: date, khu: KhuTon, nhom: Nhom, san_pham_id: int | None, q: str | None
) -> Select[Any]:
    _kiem_tra_ky(tu_ngay, den_ngay)
    t0, t1 = dau_ngay_vn(tu_ngay), dau_ngay_vn(den_ngay + timedelta(days=1))
    bd = BienDongKho
    trong_ky = (bd.created_at >= t0) & (bd.created_at < t1)
    tong = lambda dieu_kien: func.coalesce(func.sum(bd.so_luong).filter(dieu_kien), 0)  # noqa: E731
    cot_loai = [tong(trong_ky & (bd.loai == loai)).label(ten) for ten, loai in LOAI_THEO_COT]
    khoa = [SanPham.id, SanPham.sku, SanPham.ten]
    cot_khoa = [SanPham.id.label("san_pham_id"), SanPham.sku, SanPham.ten]
    if nhom == Nhom.lo:
        khoa += [LoThuoc.id, LoThuoc.so_lo, LoThuoc.han_su_dung]
        cot_khoa += [LoThuoc.id.label("lo_id"), LoThuoc.so_lo, LoThuoc.han_su_dung]
    stmt = (
        select(
            *cot_khoa, tong(bd.created_at < t0).label("ton_dau"), *cot_loai, tong(bd.created_at < t1).label("ton_cuoi")
        )
        .select_from(bd)
        .join(LoThuoc, LoThuoc.id == bd.lo_id)
        .join(SanPham, SanPham.id == LoThuoc.san_pham_id)
        .where(bd.khu == khu, bd.created_at < t1)
        .group_by(*khoa)
        # Chỉ dòng có tồn đầu hoặc có phát sinh trong kỳ
        .having(or_(tong(bd.created_at < t0) != 0, func.count().filter(trong_ky) > 0))
    )
    if san_pham_id is not None:
        stmt = stmt.where(SanPham.id == san_pham_id)
    if q and q.strip():
        mau = mau_tim_like(q)
        stmt = stmt.where(
            or_(
                func.f_unaccent(func.lower(SanPham.ten)).like(func.f_unaccent(func.lower(mau)), escape="\\"),
                func.lower(SanPham.sku).like(func.lower(mau), escape="\\"),
            )
        )
    return stmt.order_by(SanPham.ten, SanPham.id, *([LoThuoc.han_su_dung, LoThuoc.id] if nhom == Nhom.lo else []))


def _dong(d) -> DongNxt:
    lo = {"lo_id": d.lo_id, "so_lo": d.so_lo, "han_su_dung": d.han_su_dung} if hasattr(d, "so_lo") else {}
    return DongNxt(
        san_pham_id=d.san_pham_id,
        sku=d.sku,
        ten=d.ten,
        ton_dau=d.ton_dau,
        ton_cuoi=d.ton_cuoi,
        **lo,
        **{ten: getattr(d, ten) for ten, _ in LOAI_THEO_COT},
    )


async def bao_cao(
    db: AsyncSession,
    tp: ThamSoPhanTrang,
    *,
    tu_ngay: date,
    den_ngay: date,
    khu: KhuTon = KhuTon.kho_ban,
    nhom: Nhom = Nhom.san_pham,
    san_pham_id: int | None = None,
    q: str | None = None,
) -> Trang[DongNxt]:
    stmt = _truy_van(tu_ngay=tu_ngay, den_ngay=den_ngay, khu=khu, nhom=nhom, san_pham_id=san_pham_id, q=q)
    dong, tong = await lay_trang(db, stmt, tp, nhieu_cot=True)
    return tao_trang([_dong(d) for d in dong], tong, tp)


def _an_toan_excel(gia_tri: str | None) -> str | None:
    """Chuỗi bắt đầu bằng `= + - @` bị Excel hiểu là công thức: thêm dấu nháy để thành văn bản."""
    if gia_tri and gia_tri[0] in "=+-@\t\r":
        return "'" + gia_tri
    return gia_tri


async def xuat_excel(
    db: AsyncSession,
    *,
    tu_ngay: date,
    den_ngay: date,
    khu: KhuTon = KhuTon.kho_ban,
    nhom: Nhom = Nhom.san_pham,
    san_pham_id: int | None = None,
    q: str | None = None,
) -> bytes:
    stmt = _truy_van(tu_ngay=tu_ngay, den_ngay=den_ngay, khu=khu, nhom=nhom, san_pham_id=san_pham_id, q=q)
    dong = (await db.execute(stmt.limit(TOI_DA_DONG_XUAT_EXCEL + 1))).all()
    if len(dong) > TOI_DA_DONG_XUAT_EXCEL:
        raise LoiNghiepVu(
            "DU_LIEU_KHONG_HOP_LE",
            "Dữ liệu không hợp lệ",
            chi_tiet=[{"truong": "tu_ngay", "loi": f"Quá {TOI_DA_DONG_XUAT_EXCEL} dòng, hãy thu hẹp bộ lọc"}],
        )
    sach = Workbook()
    ws = sach.active
    ws.title = "Nhập xuất tồn"
    ws.append([f"Báo cáo nhập - xuất - tồn {tu_ngay:%d/%m/%Y} - {den_ngay:%d/%m/%Y} ({khu.value})"])
    ws["A1"].font = Font(bold=True)
    tieu_de = ["SKU", "Sản phẩm"] + (["Số lô", "Hạn sử dụng"] if nhom == Nhom.lo else [])
    tieu_de += [
        "Tồn đầu",
        "Nhập kho",
        "Hoàn nhập",
        "Nhận trả về",
        "Điều chỉnh",
        "Xuất bán",
        "Hủy",
        "Xuất trả NCC",
        "Tồn cuối",
    ]
    ws.append(tieu_de)
    for o in ws[2]:
        o.font = Font(bold=True)
    for d in map(_dong, dong):
        ws.append(
            [_an_toan_excel(d.sku), _an_toan_excel(d.ten)]
            + ([_an_toan_excel(d.so_lo), d.han_su_dung] if nhom == Nhom.lo else [])
            + [
                d.ton_dau,
                d.nhap_kho,
                d.hoan_nhap,
                d.nhan_tra_ve,
                d.dieu_chinh,
                d.xuat_ban,
                d.huy_hang,
                d.xuat_tra_ncc,
                d.ton_cuoi,
            ]
        )
    ws.freeze_panes = "A3"
    ra = BytesIO()
    sach.save(ra)
    return ra.getvalue()
