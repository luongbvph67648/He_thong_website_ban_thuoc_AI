"""KHO13: báo cáo nhập - xuất - tồn từ thẻ kho."""

from datetime import timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import text

from app.core.thoi_gian import bay_gio_utc, hom_nay_vn
from app.models.enums import KhuTon, LoaiBienDong
from app.modules.kho import ton_kho
from app.modules.kho.ton_kho import ChungTu
from tests.factories import tao_lo, tao_san_pham

URL = "/api/v1/noi-bo/bao-cao/nhap-xuat-ton"
CT = ChungTu("phieu_huy", 1, None)


async def lui_the_kho(db, lo_id, so_ngay):
    """Đẩy mọi dòng thẻ kho của lô về `so_ngay` ngày trước (test kỳ báo cáo)."""
    await db.execute(
        text("UPDATE bien_dong_kho SET created_at = :t WHERE lo_id = :l"),
        {"t": bay_gio_utc() - timedelta(days=so_ngay), "l": lo_id},
    )


def tham_so(tu, den, **them):
    return {
        "tu_ngay": str(hom_nay_vn() + timedelta(days=tu)),
        "den_ngay": str(hom_nay_vn() + timedelta(days=den)),
        **them,
    }


async def test_ton_cuoi_bang_so_luong_ton_hien_tai_khi_den_hom_nay(dang_nhap_kho, db):
    sp = await tao_san_pham(db)
    lo1 = await tao_lo(db, sp, so_luong=100)
    lo2 = await tao_lo(db, sp, so_luong=40)
    await ton_kho.tru_ton(db, lo1, 30, KhuTon.kho_ban, LoaiBienDong.huy_hang, CT)
    await ton_kho.dieu_chinh(db, lo2, -5, ChungTu("phieu_dieu_chinh", 1, None))
    r = (await dang_nhap_kho.get(URL, params=tham_so(-7, 0, san_pham_id=sp.id))).json()
    d = r["du_lieu"][0]
    assert (d["ton_dau"], d["nhap_kho"], d["huy_hang"], d["dieu_chinh"]) == (0, 140, -30, -5)
    assert d["ton_cuoi"] == lo1.so_luong_ton + lo2.so_luong_ton == 105
    theo_lo = (await dang_nhap_kho.get(URL, params=tham_so(-7, 0, san_pham_id=sp.id, nhom="lo"))).json()["du_lieu"]
    assert {x["lo_id"]: x["ton_cuoi"] for x in theo_lo} == {lo1.id: lo1.so_luong_ton, lo2.id: lo2.so_luong_ton}


async def test_ton_dau_la_tong_bien_dong_truoc_ky_va_dong_khong_phat_sinh_bi_bo(dang_nhap_kho, db):
    sp = await tao_san_pham(db)
    lo = await tao_lo(db, sp, so_luong=50)
    await lui_the_kho(db, lo.id, 10)  # nhập cách đây 10 ngày
    await ton_kho.tru_ton(db, lo, 8, KhuTon.kho_ban, LoaiBienDong.huy_hang, CT)  # hôm nay
    d = (await dang_nhap_kho.get(URL, params=tham_so(-3, 0, san_pham_id=sp.id))).json()["du_lieu"][0]
    assert (d["ton_dau"], d["nhap_kho"], d["huy_hang"], d["ton_cuoi"]) == (50, 0, -8, 42)
    # kỳ cũ trước khi nhập: không có dòng; kỳ giữa hai mốc không phát sinh nhưng còn tồn đầu → vẫn có dòng
    assert (await dang_nhap_kho.get(URL, params=tham_so(-30, -20, san_pham_id=sp.id))).json()["tong"] == 0
    d2 = (await dang_nhap_kho.get(URL, params=tham_so(-5, -4, san_pham_id=sp.id))).json()["du_lieu"][0]
    assert (d2["ton_dau"], d2["ton_cuoi"], d2["nhap_kho"]) == (50, 50, 0)
    assert (await dang_nhap_kho.get(URL, params=tham_so(-10, -9, san_pham_id=sp.id))).json()["du_lieu"][0][
        "nhap_kho"
    ] == 50


async def test_khu_biet_tru_tach_rieng_va_kiem_tra_ky(dang_nhap_kho, db):
    sp = await tao_san_pham(db)
    lo = await tao_lo(db, sp, so_luong=20)
    await ton_kho.vao_biet_tru(db, lo, 6, ChungTu("yeu_cau_doi_tra", 1, None))
    bt = (await dang_nhap_kho.get(URL, params=tham_so(-1, 0, san_pham_id=sp.id, khu="biet_tru"))).json()["du_lieu"][0]
    assert bt["ton_cuoi"] == lo.so_luong_biet_tru
    assert (await dang_nhap_kho.get(URL, params=tham_so(0, -1))).status_code == 400
    assert (await dang_nhap_kho.get(URL, params=tham_so(-5000, 0))).status_code == 400


async def test_excel_dung_so_lieu_va_chan_cong_thuc(dang_nhap_kho, db):
    sp = await tao_san_pham(db, ten="=HYPERLINK(1)")
    await tao_lo(db, sp, so_luong=12)
    r = await dang_nhap_kho.get(f"{URL}/excel", params=tham_so(-1, 0, san_pham_id=sp.id))
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    ws = load_workbook(BytesIO(r.content)).active
    dong = [c.value for c in ws[3]]
    assert dong[1] == "'=HYPERLINK(1)" and dong[2] == 0 and dong[-1] == 12 and dong[3] == 12


async def test_quyen_xem_cho_kho_admin_duoc_si_khong(clients_theo_vai_tro):
    p = tham_so(-1, 0)
    assert (await clients_theo_vai_tro["WAREHOUSE"].get(URL, params=p)).status_code == 200
    assert (await clients_theo_vai_tro["ADMIN"].get(f"{URL}/excel", params=p)).status_code == 200
    assert (await clients_theo_vai_tro["PHARMACIST"].get(URL, params=p)).status_code == 403
