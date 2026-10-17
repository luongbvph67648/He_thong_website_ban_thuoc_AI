"""T8.1: báo cáo tổng hợp (AD14) và báo cáo ca (DS14)."""

import io
from datetime import UTC, datetime, time, timedelta
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.thoi_gian import VN_TZ, bay_gio_utc, hom_nay_vn
from app.models.ban_hang import CtDonHang, LichSuTrangThai, PhanBoLo, ThanhToan, YeuCauDonThuoc
from tests.factories import tao_don, tao_khach, tao_lo, tao_nhan_vien, tao_san_pham

BASE = "/api/v1/noi-bo/bao-cao"
# Kỳ xa quá khứ để không lẫn với dữ liệu khác trong CSDL test: ngày thứ 500 đến thứ 491 trước hôm nay.
D0 = hom_nay_vn() - timedelta(days=500)
D1 = D0 + timedelta(days=9)


def luc(ngay_them: int, gio: int = 10) -> datetime:
    return datetime.combine(D0 + timedelta(days=ngay_them), time(gio), tzinfo=VN_TZ).astimezone(UTC)


async def don(db, khach, sp, so_luong, ngay_them, trang_thai="hoan_tat", gia_von=None, **kw):
    d = await tao_don(
        db, khach, trang_thai, dong=[{"san_pham": sp, "so_luong": so_luong}], created_at=luc(ngay_them), **kw
    )
    if gia_von is not None:
        ct = await db.scalar(select(CtDonHang).where(CtDonHang.don_hang_id == d.id))
        lo = await tao_lo(db, sp, so_luong=so_luong * 2)
        db.add(
            PhanBoLo(
                ct_don_hang_id=ct.id,
                lo_id=lo.id,
                so_luong_co_so=so_luong,
                gia_von=Decimal(gia_von),
                trang_thai="da_xuat",
                updated_at=luc(ngay_them),
            )
        )
        await db.flush()
    return d


async def bao_cao(client, **tham):
    r = await client.get(f"{BASE}/tong-hop", params={"tu_ngay": D0.isoformat(), "den_ngay": D1.isoformat(), **tham})
    assert r.status_code == 200, r.text
    return r.json()


async def test_doanh_thu_hoan_tien_gia_von_loi_nhuan(clients_theo_vai_tro, db):
    ad = clients_theo_vai_tro["ADMIN"]
    khach = await tao_khach(db)
    sp = await tao_san_pham(db, ten="SP báo cáo")
    a = await don(db, khach, sp, 3, 1, gia_von=10000)  # giá vốn 3 × 10.000
    b = await don(db, khach, sp, 2, 2, gia_von=4000)
    await don(db, khach, sp, 5, 3, trang_thai="da_huy")  # hủy: không tính
    await don(db, khach, sp, 5, 3, trang_thai="da_giao")  # chưa hoàn tất: không tính
    await don(db, khach, sp, 5, -30)  # hoàn tất nhưng ngoài kỳ
    db.add(
        ThanhToan(
            don_hang_id=a.id,
            loai="hoan_tien",
            cong_thanh_toan="tien_mat",
            so_tien=Decimal(1500),
            trang_thai="thanh_cong",
            created_at=luc(4),
        )
    )
    db.add(
        ThanhToan(
            don_hang_id=a.id,
            loai="hoan_tien",
            cong_thanh_toan="tien_mat",
            so_tien=Decimal(999),
            trang_thai="that_bai",
            created_at=luc(4),
        )
    )
    db.add(
        ThanhToan(
            don_hang_id=b.id,
            loai="hoan_tien",
            cong_thanh_toan="tien_mat",
            so_tien=Decimal(7),
            trang_thai="thanh_cong",
            created_at=luc(-40),
        )
    )  # ngoài kỳ
    await db.flush()
    d = await bao_cao(ad)
    doanh_thu = a.tong_tien + b.tong_tien
    assert d["so_don_hoan_tat"] == 2 and Decimal(d["doanh_thu"]) == doanh_thu
    assert Decimal(d["hoan_tien"]) == Decimal(1500) and Decimal(d["doanh_thu_thuan"]) == doanh_thu - 1500
    assert Decimal(d["gia_von"]) == Decimal(3 * 10000 + 2 * 4000)
    assert Decimal(d["loi_nhuan_gop"]) == doanh_thu - 1500 - 38000  # có thể âm: vẫn là số
    assert d["tu_ngay"] == D0.isoformat() and d["den_ngay"] == D1.isoformat()


async def test_phan_bo_da_hoan_nhap_khong_tinh_gia_von(clients_theo_vai_tro, db):
    ad = clients_theo_vai_tro["ADMIN"]
    khach = await tao_khach(db)
    sp = await tao_san_pham(db)
    d = await don(db, khach, sp, 2, 1)
    ct = await db.scalar(select(CtDonHang).where(CtDonHang.don_hang_id == d.id))
    lo = await tao_lo(db, sp, so_luong=10)
    db.add(
        PhanBoLo(
            ct_don_hang_id=ct.id,
            lo_id=lo.id,
            so_luong_co_so=2,
            gia_von=Decimal(9000),
            trang_thai="da_hoan_nhap",
            updated_at=luc(1),
        )
    )
    await db.flush()
    assert Decimal((await bao_cao(ad))["gia_von"]) == 0


async def test_moc_hoan_tat_lay_tu_lich_su_trang_thai(clients_theo_vai_tro, db):
    ad = clients_theo_vai_tro["ADMIN"]
    khach = await tao_khach(db)
    sp = await tao_san_pham(db)
    d = await don(db, khach, sp, 1, -100)  # tạo rất sớm, ngoài kỳ
    db.add(LichSuTrangThai(don_hang_id=d.id, tu_trang_thai="da_giao", den_trang_thai="hoan_tat", created_at=luc(5)))
    await db.flush()
    b = await bao_cao(ad)
    assert (
        b["so_don_hoan_tat"] == 1 and Decimal(b["doanh_thu"]) == d.tong_tien
    )  # tính theo lúc hoàn tất, không theo lúc tạo


async def test_doanh_thu_theo_ngay_dien_du_ngay_khong_ban_gio_vn(clients_theo_vai_tro, db):
    ad = clients_theo_vai_tro["ADMIN"]
    khach = await tao_khach(db)
    sp = await tao_san_pham(db)
    muon = await don(db, khach, sp, 1, 3, created_at=None) if False else None
    d1 = await don(db, khach, sp, 1, 3)
    d2 = await don(db, khach, sp, 2, 3)
    cuoi = await tao_don(
        db,
        khach,
        "hoan_tat",
        dong=[{"san_pham": sp, "so_luong": 1}],
        created_at=datetime.combine(D0 + timedelta(days=9), time(23, 30), tzinfo=VN_TZ).astimezone(UTC),
    )
    b = await bao_cao(ad)
    ngay = b["doanh_thu_theo_ngay"]
    assert len(ngay) == 10 and [x["ngay"] for x in ngay][0] == D0.isoformat() and ngay[-1]["ngay"] == D1.isoformat()
    assert ngay[3]["so_don"] == 2 and Decimal(ngay[3]["doanh_thu"]) == d1.tong_tien + d2.tong_tien
    assert (
        ngay[-1]["so_don"] == 1 and Decimal(ngay[-1]["doanh_thu"]) == cuoi.tong_tien
    )  # 23:30 giờ VN vẫn là ngày cuối kỳ
    assert ngay[0]["so_don"] == 0 and Decimal(ngay[0]["doanh_thu"]) == 0 and muon is None


async def test_ban_chay_xep_theo_so_luong_co_so_va_tinh_he_so_don_vi(clients_theo_vai_tro, db):
    ad = clients_theo_vai_tro["ADMIN"]
    khach = await tao_khach(db)
    don_vi = [
        {"ten_don_vi": "viên", "he_so": 1, "gia_ban": 1000},
        {"ten_don_vi": "hộp", "he_so": 20, "gia_ban": 18000, "la_mac_dinh": True},
    ]
    nhieu = await tao_san_pham(db, ten="Nhiều", don_vi=don_vi)
    it = await tao_san_pham(db, ten="Ít")
    await tao_don(
        db,
        khach,
        "hoan_tat",
        dong=[{"san_pham": nhieu, "so_luong": 2, "don_vi": "hộp"}, {"san_pham": it, "so_luong": 5}],
        created_at=luc(2),
    )
    ds = (await bao_cao(ad))["ban_chay"]
    assert (
        [x["san_pham_id"] for x in ds[:2]] == [nhieu.id, it.id]
        and ds[0]["so_luong_co_so"] == 40
        and ds[1]["so_luong_co_so"] == 5
    )
    assert Decimal(ds[0]["doanh_thu"]) == Decimal(36000)


async def test_ton_kho_va_can_han(clients_theo_vai_tro, db):
    ad = clients_theo_vai_tro["ADMIN"]
    truoc = await bao_cao(ad)
    sp = await tao_san_pham(db, ten="Tồn cận hạn")
    await tao_lo(
        db, sp, so_luong=10, gia_nhap=Decimal(2000), han_su_dung=hom_nay_vn() + timedelta(days=30), trang_thai="can_han"
    )
    await tao_lo(db, sp, so_luong=7, gia_nhap=Decimal(1000), trang_thai="dang_ban")
    await tao_lo(
        db, sp, so_luong=5, gia_nhap=Decimal(3000), han_su_dung=hom_nay_vn() - timedelta(days=1), trang_thai="het_han"
    )
    sau = await bao_cao(ad)
    assert Decimal(sau["gia_tri_ton_kho"]) - Decimal(truoc["gia_tri_ton_kho"]) == Decimal(
        10 * 2000 + 7 * 1000
    )  # lô hết hạn không tính
    theo = {x["trang_thai"]: x for x in sau["ton_kho_theo_trang_thai"]}
    het = {x["trang_thai"]: x for x in truoc["ton_kho_theo_trang_thai"]}.get("het_han", {"so_luong": 0})
    assert theo["het_han"]["so_luong"] - het["so_luong"] == 5
    assert sau["so_lo_can_han"] - truoc["so_lo_can_han"] == 1
    lo = (
        next(x for x in sau["lo_can_han"] if x["ten_san_pham"] == "Tồn cận hạn") if sau["so_lo_can_han"] <= 20 else None
    )
    if lo:
        assert lo["con_ngay"] == 30 and lo["so_luong_ton"] == 10 and Decimal(lo["gia_tri"]) == 20000


async def test_hieu_qua_ai_co_trong_bao_cao(clients_theo_vai_tro):
    b = await bao_cao(clients_theo_vai_tro["ADMIN"])
    assert {"tong_hoi_thoai", "ty_le_chuyen_duoc_si", "tong_chi_phi_usd", "danh_gia"} <= set(b["hieu_qua_ai"])


async def test_ky_khong_hop_le(clients_theo_vai_tro):
    ad = clients_theo_vai_tro["ADMIN"]
    r = await ad.get(f"{BASE}/tong-hop", params={"tu_ngay": "2026-10-10", "den_ngay": "2026-10-01"})
    assert r.status_code == 400 and r.json()["chi_tiet"][0]["truong"] == "tu_ngay"
    r = await ad.get(f"{BASE}/tong-hop", params={"tu_ngay": "2024-01-01", "den_ngay": "2026-01-01"})
    assert r.status_code == 400 and r.json()["chi_tiet"][0]["truong"] == "den_ngay"
    assert (await ad.get(f"{BASE}/tong-hop")).status_code == 400


async def test_xuat_excel_nhieu_sheet_va_chong_cong_thuc(clients_theo_vai_tro, db):
    ad = clients_theo_vai_tro["ADMIN"]
    khach = await tao_khach(db)
    sp = await tao_san_pham(db, ten='=HYPERLINK("http://x")', sku="@SKU-EX")
    await don(db, khach, sp, 1, 1)
    r = await ad.get(f"{BASE}/tong-hop/excel", params={"tu_ngay": D0.isoformat(), "den_ngay": D1.isoformat()})
    assert (
        r.status_code == 200
        and "spreadsheetml" in r.headers["content-type"]
        and "attachment" in r.headers["content-disposition"]
    )
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Tổng quan", "Theo ngày", "Bán chạy", "Cận hạn"]
    ban_chay = [[c.value for c in row] for row in wb["Bán chạy"].iter_rows()]
    assert ban_chay[1][0] == "'@SKU-EX" and ban_chay[1][1].startswith("'=HYPERLINK")  # không thành công thức
    assert len(list(wb["Theo ngày"].iter_rows())) == 11
    tong = {row[0].value: row[1].value for row in wb["Tổng quan"].iter_rows(min_row=2)}
    assert tong["Số đơn hoàn tất"] == 1


async def test_phan_quyen_bao_cao_tong_hop(clients_theo_vai_tro):
    for v in ("PHARMACIST", "WAREHOUSE"):
        assert (
            await clients_theo_vai_tro[v].get(
                f"{BASE}/tong-hop", params={"tu_ngay": D0.isoformat(), "den_ngay": D1.isoformat()}
            )
        ).status_code == 403
        assert (
            await clients_theo_vai_tro[v].get(
                f"{BASE}/tong-hop/excel", params={"tu_ngay": D0.isoformat(), "den_ngay": D1.isoformat()}
            )
        ).status_code == 403


# --- DS14 ---


def ca(tu_them=0, den_them=9, **kw):
    return {
        "tu": luc(tu_them, 0).astimezone(VN_TZ).replace(tzinfo=None).isoformat(),
        "den": luc(den_them, 23).astimezone(VN_TZ).replace(tzinfo=None).isoformat(),
        **kw,
    }


async def test_ds14_chi_tinh_viec_cua_chinh_minh_khong_co_gia_von(dang_nhap_duoc_si, db):
    me = dang_nhap_duoc_si.tai_khoan
    khac = await tao_nhan_vien(db, "PHARMACIST")
    khach = await tao_khach(db)
    sp = await tao_san_pham(db)
    tao_cua_toi = await don(db, khach, sp, 2, 1, trang_thai="hoan_tat", nhan_vien_tao_id=me.id)
    xac_nhan = await don(db, khach, sp, 1, 2, trang_thai="da_xac_nhan")
    db.add(
        LichSuTrangThai(
            don_hang_id=xac_nhan.id,
            tu_trang_thai="cho_xac_nhan",
            den_trang_thai="da_xac_nhan",
            nhan_vien_id=me.id,
            created_at=luc(2),
        )
    )
    cua_khac = await don(db, khach, sp, 9, 1, trang_thai="hoan_tat", nhan_vien_tao_id=khac.id)
    huy = await don(db, khach, sp, 1, 1, trang_thai="da_huy", nhan_vien_tao_id=me.id)
    ngoai_ca = await don(db, khach, sp, 1, -20, trang_thai="hoan_tat", nhan_vien_tao_id=me.id)
    db.add(
        YeuCauDonThuoc(
            ma_yeu_cau="YCDS14A",
            ho_ten="A",
            so_dien_thoai="0912345678",
            anh_don_thuoc=["x.png"],
            trang_thai="da_tao_don",
            duoc_si_id=me.id,
            updated_at=luc(3),
        )
    )
    db.add(
        YeuCauDonThuoc(
            ma_yeu_cau="YCDS14B",
            ho_ten="B",
            so_dien_thoai="0912345678",
            anh_don_thuoc=["x.png"],
            trang_thai="da_tao_don",
            duoc_si_id=khac.id,
            updated_at=luc(3),
        )
    )
    db.add(
        YeuCauDonThuoc(
            ma_yeu_cau="YCDS14C",
            ho_ten="C",
            so_dien_thoai="0912345678",
            anh_don_thuoc=["x.png"],
            trang_thai="tu_choi",
            duoc_si_id=me.id,
            updated_at=luc(3),
        )
    )
    for nv, so_tien, cong, ngay in (
        (me, 50000, "tien_mat", 1),
        (me, 70000, "tien_mat", 2),
        (me, 30000, "chuyen_khoan", 2),
        (khac, 99000, "tien_mat", 1),
        (me, 11000, "tien_mat", -20),
    ):
        db.add(
            ThanhToan(
                don_hang_id=tao_cua_toi.id,
                loai="thanh_toan",
                cong_thanh_toan=cong,
                so_tien=Decimal(so_tien),
                trang_thai="thanh_cong",
                nhan_vien_xu_ly_id=nv.id,
                created_at=luc(ngay),
            )
        )
    db.add(
        ThanhToan(
            don_hang_id=tao_cua_toi.id,
            loai="thanh_toan",
            cong_thanh_toan="tien_mat",
            so_tien=Decimal(5000),
            trang_thai="that_bai",
            nhan_vien_xu_ly_id=me.id,
            created_at=luc(1),
        )
    )
    await db.flush()
    r = await dang_nhap_duoc_si.get(f"{BASE}/ca", params=ca())
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["so_don"] == 2 and Decimal(d["doanh_thu"]) == tao_cua_toi.tong_tien + xac_nhan.tong_tien
    assert d["so_don_thuoc_da_duyet"] == 1
    assert (
        Decimal(d["tien_mat_thu_tai_quay"]) == 120000
        and d["so_giao_dich_tien_mat"] == 2
        and Decimal(d["chuyen_khoan_thu_tai_quay"]) == 30000
    )
    assert not {"gia_von", "loi_nhuan_gop", "loi_nhuan"} & set(d)
    assert cua_khac.id and huy.id and ngoai_ca.id


async def test_ds14_gio_khong_mui_gio_la_gio_viet_nam_va_gioi_han(dang_nhap_duoc_si):
    r = await dang_nhap_duoc_si.get(f"{BASE}/ca", params={"tu": "2026-01-01T08:00:00", "den": "2026-01-01T17:00:00"})
    assert r.status_code == 200 and r.json()["tu"].endswith("+07:00")
    dao = await dang_nhap_duoc_si.get(f"{BASE}/ca", params={"tu": "2026-01-02T08:00:00", "den": "2026-01-01T17:00:00"})
    assert dao.status_code == 400 and dao.json()["chi_tiet"][0]["truong"] == "tu"
    dai = await dang_nhap_duoc_si.get(f"{BASE}/ca", params={"tu": "2026-01-01T00:00:00", "den": "2026-03-01T00:00:00"})
    assert dai.status_code == 400 and dai.json()["chi_tiet"][0]["truong"] == "den"


async def test_ds14_chi_duoc_si(clients_theo_vai_tro):
    for v in ("ADMIN", "WAREHOUSE"):
        assert (await clients_theo_vai_tro[v].get(f"{BASE}/ca", params=ca())).status_code == 403
    assert bay_gio_utc()
